"""Test export dedup and cleanup-duplicates endpoint."""
import os
import io
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://frozen-cols-fix.preview.emergentagent.com").rstrip("/")
ADMIN_EMAIL = "admin@mhpfintech.com"
ADMIN_PASS = "Admin@123"

MONTH_KEY = "Sep-2026"
TEST_CUSTOMER = "TEST_DEDUP_CUST"
TEST_CUSTOMER_DIFF = "TEST_DEDUP_DIFFSTATUS"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASS})
    assert r.status_code == 200, r.text
    return r.json()["access_token"] if "access_token" in r.json() else r.json().get("token")


@pytest.fixture(scope="module")
def headers(token):
    return {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}


def _cleanup(headers):
    r = requests.get(f"{BASE_URL}/api/loans", headers=headers, params={"limit": 100000})
    if r.status_code != 200:
        return
    data = r.json()
    loans = data if isinstance(data, list) else data.get("loans", data.get("items", []))
    for ln in loans:
        cust = str(ln.get("customer_name", ""))
        gm = str(ln.get("group_month", ""))
        m = str(ln.get("month", ""))
        if cust.startswith("TEST_DEDUP") or (gm == MONTH_KEY and cust.startswith("TEST_")):
            requests.delete(f"{BASE_URL}/api/loans/{ln['id']}", headers=headers)


@pytest.fixture(scope="module", autouse=True)
def cleanup_before_and_after(headers):
    _cleanup(headers)
    yield
    _cleanup(headers)


def _make_loan(customer=TEST_CUSTOMER, status="Pending"):
    return {
        "agent_name": "TestAgent",
        "customer_name": customer,
        "contact_no": "9999999999",
        "company_name": "TestCo",
        "bank": "SBI",
        "status": status,
        "sanction": "100000",
        "disbursed": "50000",
        "month": MONTH_KEY,
        "group_month": MONTH_KEY,
    }


def test_create_4_identical_loans(headers):
    ids = []
    for _ in range(4):
        r = requests.post(f"{BASE_URL}/api/loans", headers=headers, json=_make_loan())
        assert r.status_code in (200, 201), f"Create failed: {r.status_code} {r.text}"
        ids.append(r.json().get("id"))
    assert len(ids) == 4
    assert len(set(ids)) == 4  # 4 distinct IDs stored


def test_export_month_dedupes(headers):
    r = requests.get(f"{BASE_URL}/api/backup/export-month/{MONTH_KEY}", headers=headers)
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data_rows = rows[1:]
    # Find customer name column index
    cust_idx = header.index("Customer Name")
    matching = [row for row in data_rows if row[cust_idx] == TEST_CUSTOMER]
    assert len(matching) == 1, f"Expected 1 deduped row, got {len(matching)}: {matching}"


def test_cleanup_duplicates_endpoint(headers):
    r = requests.post(f"{BASE_URL}/api/loans/cleanup-duplicates", headers=headers)
    assert r.status_code == 200, r.text
    data = r.json()
    assert "duplicates_removed" in data
    # Should have removed the 3 duplicate TEST_DEDUP_CUST loans (plus possibly other pre-existing dups)
    assert data["duplicates_removed"] >= 3, f"Expected >=3 removed, got {data}"

    # Verify only 1 loan with TEST_DEDUP_CUST remains
    r2 = requests.get(f"{BASE_URL}/api/loans", headers=headers, params={"limit": 100000})
    assert r2.status_code == 200
    body = r2.json()
    loans = body if isinstance(body, list) else body.get("loans", body.get("items", []))
    remaining = [l for l in loans if l.get("customer_name") == TEST_CUSTOMER]
    assert len(remaining) == 1, f"Expected 1 remaining, got {len(remaining)}"


def test_different_status_not_deduped(headers):
    # Create 2 loans same customer/contact/bank but different status
    r1 = requests.post(f"{BASE_URL}/api/loans", headers=headers,
                       json=_make_loan(customer=TEST_CUSTOMER_DIFF, status="Pending"))
    r2 = requests.post(f"{BASE_URL}/api/loans", headers=headers,
                       json=_make_loan(customer=TEST_CUSTOMER_DIFF, status="Approved"))
    assert r1.status_code in (200, 201)
    assert r2.status_code in (200, 201)

    r = requests.get(f"{BASE_URL}/api/backup/export-month/{MONTH_KEY}", headers=headers)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    cust_idx = header.index("Customer Name")
    matching = [row for row in rows[1:] if row[cust_idx] == TEST_CUSTOMER_DIFF]
    assert len(matching) == 2, f"Expected 2 rows (different status), got {len(matching)}: {matching}"


def test_full_export_dedup(headers):
    # Full export should also dedup. Create a duplicate again for TEST_DEDUP_CUST
    for _ in range(2):
        requests.post(f"{BASE_URL}/api/loans", headers=headers, json=_make_loan())

    r = requests.get(f"{BASE_URL}/api/export/loans", headers=headers, params={"month": "all"})
    assert r.status_code == 200, r.text
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    cust_idx = header.index("Customer Name")
    matching = [row for row in rows[1:] if row[cust_idx] == TEST_CUSTOMER]
    assert len(matching) == 1, f"Full export expected 1 deduped row, got {len(matching)}"
