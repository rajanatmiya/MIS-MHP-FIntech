"""Backend tests for month export filter (group_month primary + regex fallback) and cleanup-duplicates."""
import os
import pytest
import requests
from io import BytesIO
from openpyxl import load_workbook

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://frozen-cols-fix.preview.emergentagent.com').rstrip('/')
ADMIN_EMAIL = "admin@mhpfintech.com"
ADMIN_PASSWORD = "Admin@123"

TEST_PREFIX = "TEST_MONTHEXP_"


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, f"Login failed: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}", "Content-Type": "application/json"}


def _create_loan(headers, customer_name, month, group_month, bank="HDFC", status="Pending"):
    payload = {
        "agent_name": "TEST_AGENT",
        "customer_name": customer_name,
        "company_name": "TEST_CO",
        "contact_no": "9999999999",
        "status": status,
        "bank": bank,
        "month": month,
        "group_month": group_month,
    }
    r = requests.post(f"{BASE_URL}/api/loans", json=payload, headers=headers)
    assert r.status_code == 200, f"Create failed: {r.status_code} {r.text}"
    return r.json()


def _read_excel_customer_names(content_bytes):
    wb = load_workbook(BytesIO(content_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers_row = list(rows[0])
    try:
        cust_idx = headers_row.index("Customer Name")
    except ValueError:
        return []
    return [row[cust_idx] for row in rows[1:] if row[cust_idx]]


@pytest.fixture(scope="module")
def created_loans(headers):
    """Create test loans across months. Cleanup after tests."""
    created = []

    # 2 loans in Apr-2026 (group_month=Apr-2026)
    for i in range(2):
        loan = _create_loan(headers, f"{TEST_PREFIX}APR_{i}", "15-04-2026", "Apr-2026")
        created.append(loan["id"])

    # 2 loans in Aug-2026 (group_month=Aug-2026)
    for i in range(2):
        loan = _create_loan(headers, f"{TEST_PREFIX}AUG_{i}", "10-08-2026", "Aug-2026")
        created.append(loan["id"])

    # Loan with month=01-07-2026 but group_month=Aug-2026 (should appear in Aug, not Jul)
    loan = _create_loan(headers, f"{TEST_PREFIX}CROSS_JUL_INTO_AUG", "01-07-2026", "Aug-2026")
    created.append(loan["id"])
    cross_id = loan["id"]

    # Loan with group_month='' and month='15-04-2026' -> fallback regex should include in Apr
    loan = _create_loan(headers, f"{TEST_PREFIX}FALLBACK_APR", "15-04-2026", "")
    created.append(loan["id"])
    fallback_id = loan["id"]

    yield {"all_ids": created, "cross_id": cross_id, "fallback_id": fallback_id}

    # Cleanup
    for lid in created:
        try:
            requests.delete(f"{BASE_URL}/api/loans/{lid}", headers=headers)
        except Exception:
            pass


class TestMonthExport:
    """Verify month export uses group_month as authoritative filter."""

    def test_export_apr_2026_contains_only_april(self, headers, created_loans):
        r = requests.get(f"{BASE_URL}/api/backup/export-month/Apr-2026", headers=headers)
        assert r.status_code == 200, f"Export failed: {r.status_code} {r.text}"
        names = _read_excel_customer_names(r.content)
        test_names = [n for n in names if n and str(n).startswith(TEST_PREFIX)]
        # Expect APR_0, APR_1, and FALLBACK_APR
        assert f"{TEST_PREFIX}APR_0" in test_names, f"Missing APR_0. Got: {test_names}"
        assert f"{TEST_PREFIX}APR_1" in test_names, f"Missing APR_1. Got: {test_names}"
        assert f"{TEST_PREFIX}FALLBACK_APR" in test_names, f"Fallback regex failed. Got: {test_names}"
        # Ensure NO Aug entries leak in
        aug_leaked = [n for n in test_names if "AUG_" in n or "CROSS_JUL_INTO_AUG" in n]
        assert not aug_leaked, f"Aug entries leaked into Apr export: {aug_leaked}"

    def test_export_aug_2026_contains_only_august(self, headers, created_loans):
        r = requests.get(f"{BASE_URL}/api/backup/export-month/Aug-2026", headers=headers)
        assert r.status_code == 200
        names = _read_excel_customer_names(r.content)
        test_names = [n for n in names if n and str(n).startswith(TEST_PREFIX)]
        assert f"{TEST_PREFIX}AUG_0" in test_names
        assert f"{TEST_PREFIX}AUG_1" in test_names
        # Cross-month loan (month=Jul, group_month=Aug) should be here
        assert f"{TEST_PREFIX}CROSS_JUL_INTO_AUG" in test_names, \
            f"group_month=Aug-2026 not authoritative. Got: {test_names}"
        # No Apr entries should leak
        apr_leaked = [n for n in test_names if "APR_" in n or "FALLBACK_APR" in n]
        assert not apr_leaked, f"Apr entries leaked into Aug: {apr_leaked}"

    def test_export_jul_2026_excludes_cross_loan(self, headers, created_loans):
        """Loan with month=01-07-2026 but group_month=Aug-2026 must NOT appear in Jul export."""
        r = requests.get(f"{BASE_URL}/api/backup/export-month/Jul-2026", headers=headers)
        assert r.status_code == 200
        names = _read_excel_customer_names(r.content)
        test_names = [n for n in names if n and str(n).startswith(TEST_PREFIX)]
        assert f"{TEST_PREFIX}CROSS_JUL_INTO_AUG" not in test_names, \
            f"Cross-month loan leaked into Jul because month field says Jul. Got: {test_names}"


class TestCleanupDuplicates:
    """Verify cleanup-duplicates endpoint."""

    def test_cleanup_removes_duplicates(self, headers):
        # Create 3 identical duplicates
        dup_ids = []
        for i in range(3):
            payload = {
                "agent_name": "TEST_AGENT",
                "customer_name": f"{TEST_PREFIX}DUP_CUST",
                "company_name": "TEST_CO",
                "contact_no": "8888888888",
                "status": "Pending",
                "bank": "ICICI",
                "month": "15-05-2026",
                "group_month": "May-2026",
            }
            r = requests.post(f"{BASE_URL}/api/loans", json=payload, headers=headers)
            assert r.status_code == 200
            dup_ids.append(r.json()["id"])

        # Count before via listing
        r = requests.get(f"{BASE_URL}/api/loans?search={TEST_PREFIX}DUP_CUST&limit=100", headers=headers)
        assert r.status_code == 200
        before_count = len([l for l in r.json()["loans"] if l.get("customer_name") == f"{TEST_PREFIX}DUP_CUST"])
        assert before_count == 3, f"Expected 3 duplicates, got {before_count}"

        # Run cleanup
        r = requests.post(f"{BASE_URL}/api/loans/cleanup-duplicates", headers=headers)
        assert r.status_code == 200
        result = r.json()
        assert "duplicates_removed" in result
        assert result["duplicates_removed"] >= 2, \
            f"Expected >=2 duplicates removed, got {result['duplicates_removed']}"

        # Verify only 1 remains for our test customer
        r = requests.get(f"{BASE_URL}/api/loans?search={TEST_PREFIX}DUP_CUST&limit=100", headers=headers)
        assert r.status_code == 200
        remaining = [l for l in r.json()["loans"] if l.get("customer_name") == f"{TEST_PREFIX}DUP_CUST"]
        assert len(remaining) == 1, f"Expected 1 remaining, got {len(remaining)}"

        # Cleanup remaining
        for l in remaining:
            requests.delete(f"{BASE_URL}/api/loans/{l['id']}", headers=headers)


@pytest.fixture(scope="module", autouse=True)
def final_cleanup(headers):
    """After all tests, scan and delete any leftover TEST_MONTHEXP_ loans."""
    yield
    r = requests.get(f"{BASE_URL}/api/loans?search={TEST_PREFIX}&limit=500", headers=headers)
    if r.status_code == 200:
        for loan in r.json().get("loans", []):
            if str(loan.get("customer_name", "")).startswith(TEST_PREFIX):
                try:
                    requests.delete(f"{BASE_URL}/api/loans/{loan['id']}", headers=headers)
                except Exception:
                    pass
